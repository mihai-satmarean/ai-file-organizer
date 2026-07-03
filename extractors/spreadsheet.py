# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0
# Required Notice: Copyright (c) 2026 George Zhang — https://github.com/TheYellowDuck

"""
Spreadsheet content extractor for Excel files.
"""

from pathlib import Path
from .base import BaseExtractor, ExtractedContent

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class SpreadsheetExtractor(BaseExtractor):
    """Extract content from Excel spreadsheets."""
    
    SUPPORTED_EXTENSIONS = {'.xlsx', '.xlsm', '.xltx', '.xltm'}
    SUPPORTED_MIME_TYPES = {
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    }
    
    # Maximum rows to extract per sheet
    MAX_ROWS = 500
    # Maximum sheets to process
    MAX_SHEETS = 10
    
    def extract(self, file_path: Path) -> ExtractedContent:
        """Extract text content from an Excel file."""
        if not OPENPYXL_AVAILABLE:
            return self._create_error_result(
                file_path,
                "openpyxl not installed. Run: pip install openpyxl"
            )
        
        try:
            wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
            
            metadata = {
                "sheets": wb.sheetnames[:self.MAX_SHEETS],
                "total_sheets": len(wb.sheetnames)
            }
            
            all_text_parts = []
            
            for sheet_name in wb.sheetnames[:self.MAX_SHEETS]:
                sheet = wb[sheet_name]
                sheet_text = [f"[Sheet: {sheet_name}]"]
                
                row_count = 0
                for row in sheet.iter_rows(max_row=self.MAX_ROWS, values_only=True):
                    # Filter out None values and convert to strings
                    row_values = [str(cell) for cell in row if cell is not None]
                    if row_values:
                        sheet_text.append(' | '.join(row_values))
                        row_count += 1
                
                if row_count > 0:
                    all_text_parts.append('\n'.join(sheet_text))
                    metadata[f"rows_{sheet_name}"] = row_count
            
            wb.close()
            
            text_content = '\n\n'.join(all_text_parts)
            
            return self._create_success_result(
                file_path,
                text_content,
                metadata
            )
            
        except Exception as e:
            return self._create_error_result(file_path, f"Excel extraction error: {str(e)}")
