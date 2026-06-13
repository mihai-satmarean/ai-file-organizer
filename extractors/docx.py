"""
Word document content extractor using python-docx.
"""

from pathlib import Path
from .base import BaseExtractor, ExtractedContent

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class DocxExtractor(BaseExtractor):
    """Extract text content from Word documents (.docx)."""
    
    SUPPORTED_EXTENSIONS = {'.docx'}
    SUPPORTED_MIME_TYPES = {
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    }
    
    def extract(self, file_path: Path) -> ExtractedContent:
        """Extract text content from a Word document."""
        if not DOCX_AVAILABLE:
            return self._create_error_result(
                file_path,
                "python-docx not installed. Run: pip install python-docx"
            )
        
        try:
            doc = Document(str(file_path))
            
            # Extract core properties/metadata
            metadata = {}
            try:
                props = doc.core_properties
                metadata = {
                    "title": props.title or "",
                    "author": props.author or "",
                    "subject": props.subject or "",
                    "keywords": props.keywords or "",
                    "created": str(props.created) if props.created else "",
                    "modified": str(props.modified) if props.modified else "",
                }
                metadata = {k: v for k, v in metadata.items() if v}
            except Exception:
                pass
            
            # Extract paragraphs
            paragraphs = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)
            
            # Extract text from tables
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        table_texts.append(' | '.join(row_text))
            
            # Combine all text
            all_text_parts = []
            
            # Add title if available
            if metadata.get('title'):
                all_text_parts.append(f"Title: {metadata['title']}")
            
            if paragraphs:
                all_text_parts.extend(paragraphs)
            
            if table_texts:
                all_text_parts.append("\n[Tables]")
                all_text_parts.extend(table_texts)
            
            text_content = '\n\n'.join(all_text_parts)
            
            metadata["paragraphs"] = len(paragraphs)
            metadata["tables"] = len(doc.tables)
            
            return self._create_success_result(
                file_path,
                text_content,
                metadata
            )
            
        except Exception as e:
            return self._create_error_result(file_path, f"DOCX extraction error: {str(e)}")
