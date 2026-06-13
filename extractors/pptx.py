"""
PowerPoint content extractor.
"""

from pathlib import Path
from .base import BaseExtractor, ExtractedContent

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


class PowerPointExtractor(BaseExtractor):
    """Extract content from PowerPoint presentations."""
    
    SUPPORTED_EXTENSIONS = {'.pptx', '.pptm'}
    SUPPORTED_MIME_TYPES = {
        'application/vnd.openxmlformats-officedocument.presentationml.presentation'
    }
    
    def extract(self, file_path: Path) -> ExtractedContent:
        """Extract text content from a PowerPoint file."""
        if not PPTX_AVAILABLE:
            return self._create_error_result(
                file_path,
                "python-pptx not installed. Run: pip install python-pptx"
            )
        
        try:
            prs = Presentation(str(file_path))
            
            metadata = {
                "slides": len(prs.slides)
            }
            
            # Extract core properties
            try:
                if prs.core_properties:
                    props = prs.core_properties
                    if props.title:
                        metadata["title"] = props.title
                    if props.author:
                        metadata["author"] = props.author
                    if props.subject:
                        metadata["subject"] = props.subject
            except Exception:
                pass
            
            all_text_parts = []
            
            # Add title if available
            if metadata.get("title"):
                all_text_parts.append(f"Title: {metadata['title']}")
            
            for i, slide in enumerate(prs.slides, 1):
                slide_text = [f"[Slide {i}]"]
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                    
                    # Extract text from tables
                    if shape.has_table:
                        table = shape.table
                        for row in table.rows:
                            row_text = []
                            for cell in row.cells:
                                if cell.text.strip():
                                    row_text.append(cell.text.strip())
                            if row_text:
                                slide_text.append(' | '.join(row_text))
                
                if len(slide_text) > 1:  # More than just the slide header
                    all_text_parts.append('\n'.join(slide_text))
            
            text_content = '\n\n'.join(all_text_parts)
            
            return self._create_success_result(
                file_path,
                text_content,
                metadata
            )
            
        except Exception as e:
            return self._create_error_result(file_path, f"PowerPoint extraction error: {str(e)}")
